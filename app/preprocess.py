"""Step 0 — no LLM: text extraction, input validation and PII redaction.

Everything here runs locally *before* any text is sent to the LLM API.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

MIN_CV_CHARS = 300
MIN_JOB_CHARS = 150
MAX_CHARS = 15_000
MAX_FILE_BYTES = 4 * 1024 * 1024  # Vercel caps function request bodies at 4.5 MB


class InputError(ValueError):
    """User-facing validation error. ``code`` maps to a translated message."""

    def __init__(self, code: str, **params):
        self.code = code
        self.params = params
        super().__init__(code)


MESSAGES = {
    "en": {
        "cv_missing": "Please upload your CV or paste its text.",
        "job_missing": "Please paste the job posting.",
        "cv_too_short": "Your CV text is too short ({n} characters, minimum {min}). If you uploaded a scanned PDF, "
                        "it may contain images only — please paste the text instead.",
        "job_too_short": "The job posting is too short ({n} characters, minimum {min}). Paste the full posting.",
        "cv_too_long": "Your CV is too long ({n} characters, maximum {max}). Please shorten it or paste only the CV.",
        "job_too_long": "The job posting is too long ({n} characters, maximum {max}).",
        "file_too_big": "The file is larger than 4 MB.",
        "file_type": "Unsupported file type. Please upload a PDF, DOCX or TXT file.",
        "file_unreadable": "We could not read text from this file. Please paste the CV text instead.",
    },
    "tr": {
        "cv_missing": "Lütfen CV'nizi yükleyin veya metnini yapıştırın.",
        "job_missing": "Lütfen iş ilanını yapıştırın.",
        "cv_too_short": "CV metni çok kısa ({n} karakter, en az {min}). Taranmış bir PDF yüklediyseniz yalnızca "
                        "görsel içeriyor olabilir — lütfen metni yapıştırın.",
        "job_too_short": "İş ilanı çok kısa ({n} karakter, en az {min}). Lütfen ilanın tamamını yapıştırın.",
        "cv_too_long": "CV çok uzun ({n} karakter, en fazla {max}). Lütfen kısaltın veya yalnızca CV'yi yapıştırın.",
        "job_too_long": "İş ilanı çok uzun ({n} karakter, en fazla {max}).",
        "file_too_big": "Dosya 4 MB'tan büyük.",
        "file_type": "Desteklenmeyen dosya türü. Lütfen PDF, DOCX veya TXT yükleyin.",
        "file_unreadable": "Bu dosyadan metin okunamadı. Lütfen CV metnini yapıştırın.",
    },
}


def error_message(err: InputError, lang: str) -> str:
    template = MESSAGES.get(lang, MESSAGES["en"]).get(err.code, err.code)
    return template.format(min=MIN_CV_CHARS if "cv" in err.code else MIN_JOB_CHARS, max=MAX_CHARS, **err.params)


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise InputError("file_too_big")
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx"):
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:  # CV templates often put content in tables
                for row in table.rows:
                    cells = []
                    for cell in row.cells:
                        if cell.text not in cells:  # merged cells repeat their text
                            cells.append(cell.text)
                    parts.append(" | ".join(cells))
            text = "\n".join(parts)
        elif name.endswith(".txt") or name.endswith(".md"):
            for encoding in ("utf-8-sig", "cp1254", "latin-1"):  # cp1254 = Turkish Windows
                try:
                    text = data.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
        else:
            raise InputError("file_type")
    except InputError:
        raise
    except Exception as exc:  # corrupt / encrypted files etc.
        raise InputError("file_unreadable") from exc
    return text


def normalise(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t\f\v]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- #
# Redaction
# --------------------------------------------------------------------------- #
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
# Turkish mobiles: +90 5xx xxx xx xx, 0 5xx xxx xx xx, 05xxxxxxxxx, (0532) 123 45 67 ...
TR_MOBILE_RE = re.compile(r"(?<!\d)(?:\+?90[\s.-]?)?\(?0?5\d{2}\)?[\s.-]?\d{3}[\s.-]?\d{2}[\s.-]?\d{2}(?!\d)")
# Turkish landlines with leading 0 or +90: 0212 555 12 34, +90 (216) 555 12 34
TR_LANDLINE_RE = re.compile(r"(?<!\d)(?:\+90[\s.-]?|0)\(?[2-4]\d{2}\)?[\s.-]?\d{3}[\s.-]?\d{2}[\s.-]?\d{2}(?!\d)")
# Generic international numbers: +44 20 7946 0958, +1 (555) 123-4567
INTL_PHONE_RE = re.compile(r"(?<![\w+])\+\d{1,3}[\s.-]?\(?\d{1,4}\)?(?:[\s.-]?\d{2,4}){2,4}(?!\d)")
TCKN_RE = re.compile(r"(?<!\d)[1-9]\d{10}(?!\d)")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[\w%-]+/?", re.I)

ADDRESS_LABEL_RE = re.compile(r"^\s*(?:[-•*·▪●]\s*)?(adres|address|ev adresi|home address)\b\s*[:：]", re.I)
TR_ADDRESS_TOKENS = re.compile(r"\b(mah\.|mahallesi|mh\.|cad\.|caddesi|cd\.|sok\.|sokak|sokağı|sk\.|bulvarı|blv\.)",
                               re.I)
EN_ADDRESS_RE = re.compile(r"\b\d{1,5}\s+\w+(?:\s\w+)?\s(street|st\.|avenue|ave\.|road|rd\.|boulevard|blvd\.|lane)\b",
                           re.I)

# Personal attributes that are irrelevant to job fit. Matched on a "folded"
# version of the text where İ/I/ı are all mapped to "i" (Turkish casing).
ATTRIBUTE_LABELS: dict[str, list[str]] = {
    "date_of_birth": [r"doğum tarihi", r"doğum yeri ve tarihi", r"doğum yeri", r"date of birth", r"birth ?date",
                      r"birthday", r"dob", r"d\.o\.b\.?"],
    "age": [r"yaş", r"yaşi", r"age"],
    "marital_status": [r"medeni hal", r"medeni hali", r"medeni durum", r"medeni durumu", r"marital status"],
    "gender": [r"cinsiyet", r"cinsiyeti", r"gender", r"sex"],
    "nationality": [r"uyruk", r"uyruğu", r"uyruk[u]", r"nationality", r"citizenship", r"vatandaşlik"],
    "religion": [r"din", r"dini", r"religion"],
    "military_status": [r"askerlik", r"askerlik durumu", r"military status", r"military service"],
    "photo": [r"fotoğraf", r"vesikalik", r"photo", r"photograph"],
    "health_physical": [r"kan grubu", r"blood type", r"boy", r"kilo", r"height", r"weight", r"sağlik durumu",
                        r"health status", r"engellilik durumu", r"disability"],
    "family": [r"çocuk sayisi", r"number of children", r"children", r"eş", r"spouse"],
    "national_id": [r"t\.?c\.? kimlik no", r"tc kimlik", r"t\.c\.? no", r"tckn", r"kimlik no", r"national id"],
}
_LABEL_ALT = "|".join(f"(?P<{k}>{'|'.join(v)})" for k, v in ATTRIBUTE_LABELS.items())
# label at the start of a segment, followed by ":" (or a dash / 2+ spaces / tab)
ATTRIBUTE_RE = re.compile(rf"^\s*(?:[-•*·▪●]\s*)?(?:{_LABEL_ALT})\s*(?:[:：]|\s[-–]\s|\t|\s{{2,}})\s*\S")
PHOTO_REF_RE = re.compile(r"\[(?:photo|foto|fotoğraf|image|resim)[^\]]*\]|\bvesikalık\b", re.I)
SEGMENT_SPLIT_RE = re.compile(r"(\s[|•·]\s|\t+|\s{3,})")

NAME_SKIP_WORDS = {"cv", "curriculum vitae", "resume", "résumé", "özgeçmiş", "ozgecmis", "öz geçmiş"}

LABELS_HUMAN = {
    "en": {
        "email": "email address", "phone": "phone number", "national_id": "national ID number (TC Kimlik)",
        "address": "home address", "linkedin": "LinkedIn profile URL (contains your name)",
        "name": "your name", "date_of_birth": "date of birth", "age": "age", "marital_status": "marital status",
        "gender": "gender", "nationality": "nationality", "religion": "religion",
        "military_status": "military service status", "photo": "photo reference",
        "health_physical": "health / physical details (e.g. blood type, height)",
        "family": "family details",
    },
    "tr": {
        "email": "e-posta adresi", "phone": "telefon numarası", "national_id": "TC Kimlik numarası",
        "address": "ev adresi", "linkedin": "LinkedIn profil bağlantısı (adınızı içerir)", "name": "adınız",
        "date_of_birth": "doğum tarihi", "age": "yaş", "marital_status": "medeni hal", "gender": "cinsiyet",
        "nationality": "uyruk", "religion": "din", "military_status": "askerlik durumu",
        "photo": "fotoğraf referansı", "health_physical": "sağlık / fiziksel bilgiler (ör. kan grubu, boy)",
        "family": "aile bilgileri",
    },
}


def _fold(text: str) -> str:
    """Lower-case with Turkish-safe dotted/dotless i handling (length preserving)."""
    return text.replace("İ", "i").replace("I", "i").replace("ı", "i").lower()


@dataclass
class Redaction:
    kinds: dict[str, int] = field(default_factory=dict)

    def add(self, kind: str, n: int = 1) -> None:
        if n:
            self.kinds[kind] = self.kinds.get(kind, 0) + n


def _sub(pattern: re.Pattern, kind: str, text: str, red: Redaction) -> str:
    new, n = pattern.subn(f"[REDACTED: {kind}]", text)
    red.add(kind, n)
    return new


def _redact_attribute_lines(text: str, red: Redaction) -> str:
    out_lines = []
    for line in text.split("\n"):
        parts = SEGMENT_SPLIT_RE.split(line)  # keeps separators at odd indices
        changed = False
        for i in range(0, len(parts), 2):
            seg = parts[i]
            m = ATTRIBUTE_RE.match(_fold(seg))
            if m:
                kind = next(k for k, v in m.groupdict().items() if v)
                parts[i] = f"[REDACTED: {kind}]"
                red.add(kind)
                changed = True
            elif ADDRESS_LABEL_RE.match(seg) or (
                len(seg) < 160 and (len(TR_ADDRESS_TOKENS.findall(seg)) >= 1 and re.search(r"\d", seg)
                                    or EN_ADDRESS_RE.search(seg))
            ):
                parts[i] = "[REDACTED: address]"
                red.add("address")
                changed = True
        line = "".join(parts) if changed else line
        line, n = PHOTO_REF_RE.subn("[REDACTED: photo]", line)
        red.add("photo", n)
        out_lines.append(line)
    return "\n".join(out_lines)


def _redact_name(text: str, red: Redaction) -> str:
    """Heuristic: the first short line made only of capitalised words is the candidate's name."""
    lines = text.split("\n")
    for idx, raw in enumerate(lines[:4]):
        line = raw.strip()
        if not line:
            continue
        if _fold(line).strip(" :") in NAME_SKIP_WORDS:
            continue
        words = line.split()
        if 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-ZÇĞİÖŞÜ][\wçğıöşüÇĞİÖŞÜ'’.-]*", w) for w in words) \
                and not re.search(r"\d", line):
            lines[idx] = raw.replace(line, "[REDACTED: name]")
            red.add("name")
            # Also remove later exact repetitions of the name (e.g. in a footer).
            rest = "\n".join(lines[idx + 1 :])
            rest, n = re.subn(re.escape(line), "[REDACTED: name]", rest)
            return "\n".join(lines[: idx + 1]) + ("\n" + rest if idx + 1 < len(lines) else "")
        break
    return text


def redact(text: str) -> tuple[str, Redaction]:
    red = Redaction()
    text = _sub(EMAIL_RE, "email", text, red)
    text = _sub(LINKEDIN_RE, "linkedin", text, red)
    text = _redact_attribute_lines(text, red)   # labelled lines first (TCKN label, DOB...)
    text = _sub(TCKN_RE, "national_id", text, red)
    for pattern in (TR_MOBILE_RE, TR_LANDLINE_RE, INTL_PHONE_RE):
        text = _sub(pattern, "phone", text, red)
    text = _redact_name(text, red)
    return text, red


def describe_redactions(red: Redaction, lang: str) -> list[dict]:
    labels = LABELS_HUMAN.get(lang, LABELS_HUMAN["en"])
    return [{"kind": k, "label": labels.get(k, k), "count": n} for k, n in red.kinds.items()]


# --------------------------------------------------------------------------- #
# Entry point used by the pipeline
# --------------------------------------------------------------------------- #
@dataclass
class Prepared:
    cv_redacted: str
    job_text: str
    redactions: list[dict]
    cv_chars: int


def prepare(cv_text: str, job_text: str, lang: str = "en") -> Prepared:
    cv = normalise(cv_text)
    job = normalise(job_text)
    if not cv:
        raise InputError("cv_missing")
    if not job:
        raise InputError("job_missing")
    if len(cv) < MIN_CV_CHARS:
        raise InputError("cv_too_short", n=len(cv))
    if len(cv) > MAX_CHARS:
        raise InputError("cv_too_long", n=len(cv))
    if len(job) < MIN_JOB_CHARS:
        raise InputError("job_too_short", n=len(job))
    if len(job) > MAX_CHARS:
        raise InputError("job_too_long", n=len(job))

    cv_redacted, red = redact(cv)
    # Job postings may contain recruiter emails / phones; redact those too (not reported).
    job_redacted, _ = redact_contact_only(job)
    return Prepared(cv_redacted=cv_redacted, job_text=job_redacted,
                    redactions=describe_redactions(red, lang), cv_chars=len(cv))


def redact_contact_only(text: str) -> tuple[str, Redaction]:
    red = Redaction()
    text = _sub(EMAIL_RE, "email", text, red)
    for pattern in (TR_MOBILE_RE, TR_LANDLINE_RE, INTL_PHONE_RE):
        text = _sub(pattern, "phone", text, red)
    return text, red
